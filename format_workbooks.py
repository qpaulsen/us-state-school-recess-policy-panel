"""Apply one house style to the deposit workbooks: Arial 9 pt throughout, bold white header row on navy,
top-aligned wrapped text, frozen header row and filters on tabular sheets, no highlights or conditional formats.
Usage: python format_workbooks.py ../data/state_school_policy_panel_v1.0.xlsx ../data/summary_statistics_v1.0.xlsx
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter

BODY = Font(name='Arial', size=9)
BOLD = Font(name='Arial', size=9, bold=True)
HEAD = Font(name='Arial', size=9, bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='1F3864')
NOFILL = PatternFill(fill_type=None)
TABULAR = {'Panel', 'AUDIT_LOG', 'AUDIT_CANDIDATES', 'Outdoor'}
WIDE = {'source_excerpt': 60, 'notes_ambiguities': 60, 'finding': 70, 'citation': 40, 'source_url': 35,
        'verbatim excerpt': 60, 'coding decision for PI': 50, 'source_file': 35}

def last_row(ws):
    for r in range(ws.max_row, 0, -1):
        if any(c.value not in (None, '') for c in ws[r]):
            return r
    return 1

def fmt(path):
    wb = openpyxl.load_workbook(path)
    for ws in wb:
        lr = last_row(ws)
        if ws.max_row > lr:
            ws.delete_rows(lr + 1, ws.max_row - lr)
        ws.conditional_formatting = ConditionalFormattingList()
        ws.data_validations.dataValidation = []
        tab = ws.title in TABULAR
        for row in ws.iter_rows():
            for c in row:
                c.fill = NOFILL
                bold = bool(c.font and c.font.b)
                c.font = BOLD if bold else BODY
                c.alignment = Alignment(wrap_text=True, vertical='top')
        if tab:
            for c in ws[1]:
                c.font = HEAD; c.fill = HFILL
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
            for i, c in enumerate(ws[1], start=1):
                ws.column_dimensions[get_column_letter(i)].width = WIDE.get(str(c.value), 16)
        else:
            ws.freeze_panes = None
        if ws.title == 'Summary':
            for row in ws.iter_rows(min_col=3, max_col=3):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = '0.0%'
            ws.column_dimensions['A'].width = 62
            for col in 'BCDEFG':
                ws.column_dimensions[col].width = 14
    wb.save(path)
    print('formatted', path)

if __name__ == '__main__':
    for p in sys.argv[1:]:
        fmt(p)
