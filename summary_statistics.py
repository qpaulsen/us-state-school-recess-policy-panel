"""Regenerate the summary statistics workbook from the panel (static values, no formulas).

Usage: python summary_statistics.py state_school_policy_panel_v1.0.xlsx summary_statistics_v1.0.xlsx
Sheets: Panel (copy of the Panel sheet), Summary (univariate counts, cross-tabs, date coverage,
rows per jurisdiction), Outdoor (rows carrying outdoor content).
"""
import sys, collections, datetime
import openpyxl
from openpyxl.styles import Font

UNI = ['topic', 'policy_status', 'outdoor_specificity', 'confidence', 'source_type', 'enforcement',
       'instrument_class', 'withholding_recess', 'frequency', 'date_basis', 'level']
CROSS = [('topic', 'policy_status'), ('outdoor_specificity', 'policy_status'), ('topic', 'confidence'),
         ('policy_status', 'enforcement'), ('instrument_class', 'policy_status'), ('source_type', 'confidence')]

def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['Panel']
    hdr = [c.value for c in ws[1]]
    n = max(i for i, h in enumerate(hdr) if h) + 1
    ci = {h: i for i, h in enumerate(hdr) if h}
    rows = [r[:n] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    val = lambda r, f: '(blank)' if r[ci[f]] in (None, '') else str(r[ci[f]])
    o = openpyxl.Workbook()
    p = o.active; p.title = 'Panel'
    p.append(hdr[:n])
    for r in rows:
        p.append(['' if v is None else v for v in r])
    s = o.create_sheet('Summary'); B = Font(bold=True)
    def head(*cells):
        s.append(list(cells)); s.cell(row=s.max_row, column=1).font = B
    N = len(rows)
    head(f'QAP summary statistics, regenerated {datetime.date.today().isoformat()} from {src}')
    s.append([f'All counts are static values computed from the Panel sheet ({N} coded rows).'])
    s.append([])
    head('Panel size'); s.append(['Coded rows (N)', N])
    s.append(['Jurisdictions', len({r[ci['state']] for r in rows})])
    s.append(['Non-silent (instrument) rows', sum(r[ci['policy_status']] != 'silent' for r in rows)])
    s.append(['Silent baseline rows', sum(r[ci['policy_status']] == 'silent' for r in rows)])
    for f in UNI:
        s.append([]); head(f, 'n', '% of N')
        c = collections.Counter(val(r, f) for r in rows)
        for k, v in c.most_common():
            s.append([k, v, v / N])
        s.append(['Total', N])
    ns = [r for r in rows if r[ci['policy_status']] != 'silent']
    pres = lambda r, f: r[ci[f]] not in (None, '')
    s.append([]); head('Date field coverage (non-silent rows only)', 'n', '% of non-silent')
    for lab, fn in [('effective_date present', lambda r: pres(r, 'effective_date')),
                    ('effective_date blank', lambda r: not pres(r, 'effective_date')),
                    ('enactment_date present', lambda r: pres(r, 'enactment_date')),
                    ('implementation_school_year present', lambda r: pres(r, 'implementation_school_year')),
                    ('amended_date present', lambda r: pres(r, 'amended_date')),
                    ('repeal_date present', lambda r: pres(r, 'repeal_date')),
                    ('effective AND enactment present', lambda r: pres(r, 'effective_date') and pres(r, 'enactment_date')),
                    ('treatment_school_year present', lambda r: pres(r, 'treatment_school_year'))]:
        k = sum(fn(r) for r in ns); s.append([lab, k, k / len(ns)])
    for a, b in CROSS:
        s.append([]); cols = sorted({val(r, b) for r in rows}); head(f'{a} x {b}', *cols, 'Total')
        for ka in sorted({val(r, a) for r in rows}):
            line = [sum(1 for r in rows if val(r, a) == ka and val(r, b) == kb) for kb in cols]
            s.append([ka] + line + [sum(line)])
        tot = [sum(1 for r in rows if val(r, b) == kb) for kb in cols]; s.append(['Total'] + tot + [sum(tot)])
    s.append([]); head('Rows per jurisdiction', 'all rows', 'silent', 'mandate', 'recommendation', 'high conf', 'medium conf')
    for st in sorted({r[ci['jurisdiction_name']] for r in rows}):
        sub = [r for r in rows if r[ci['jurisdiction_name']] == st]
        s.append([st, len(sub)] + [sum(r[ci['policy_status']] == k for r in sub) for k in ('silent', 'mandate', 'recommendation')]
                 + [sum(r[ci['confidence']] == k for r in sub) for k in ('high', 'medium')])
    s.append(['Total', N])
    od = o.create_sheet('Outdoor'); fields = ['row_id', 'state', 'topic', 'policy_status', 'outdoor_specificity', 'grade_bands', 'instrument_class', 'minutes_required', 'treatment_school_year']
    od.append(fields)
    for r in rows:
        if r[ci['outdoor_specificity']] in ('explicit_outdoor', 'outdoor_if_feasible', 'indoor_permitted') or r[ci['topic']] == 'outdoor_time':
            od.append([r[ci[f]] for f in fields])
    s.column_dimensions['A'].width = 60
    o.save(out); print(f'wrote {out}: {N} rows')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
